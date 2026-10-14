"""Revenue data sourcing for VROM slides.

Reads revenue from Quoter Excel files when available, falling back to
a computation from the domain model (rate × days with a default GM%).
"""

from __future__ import annotations

import contextlib
import warnings
from collections.abc import Iterable
from pathlib import Path
from typing import TYPE_CHECKING

from engine.domain.constants.columns import COL_RATE_DAY
from engine.domain.constants.files.outputs import out_quoter_xlsx
from engine.domain.services.rate_parser import parse_rate_day
from engine.logging import get_logger

if TYPE_CHECKING:
    from engine.domain.model.aggregates.phase import Phase

try:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    _HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    _HAS_OPENPYXL = False

logger = get_logger(__name__)

_DEFAULT_GM_PCT = 0.25  # 25 % gross margin fallback


def default_gm_pct() -> float:
    """Return the default gross-margin percentage."""
    return _DEFAULT_GM_PCT


def read_quoter_revenue(excel_dir: Path, filename: str) -> float | None:
    """Try to read total revenue from a Quoter workbook.

    Returns ``None`` if openpyxl is unavailable or the file doesn't exist.
    """
    if not _HAS_OPENPYXL:
        return None
    path = excel_dir / filename
    if not path.exists():
        return None
    try:
        # Try data_only first for cached formula results
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            wb = load_workbook(path, data_only=True)
        val = wb.active["B2"].value
        if val is not None and isinstance(val, (int, float)):
            return float(val)
        # Fallback: compute from raw data
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            wb = load_workbook(path)
        ws = wb.active
        gm = ws["T5"].value or 0.0
        total_cost = 0.0
        for row in range(7, 200):
            rate = ws.cell(row=row, column=4).value
            if rate is None or not isinstance(rate, (int, float)):
                continue
            days = sum(ws.cell(row=row, column=c).value or 0 for c in range(6, 12))
            total_cost += rate * days
        if isinstance(gm, (int, float)) and 0 < gm < 1:
            return total_cost / (1 - gm)
        return total_cost
    except Exception:  # noqa: BLE001
        logger.warning("Failed to read Quoter workbook %s", path, exc_info=True)
        return None


def read_gm_label(excel_dir: Path, phase_ids: Iterable[str]) -> str:
    """Try to extract the GM% from any existing Quoter file for display.

    Iterates the project's actual phases rather than a fixed phase tuple, so a
    project with any phase set finds a workbook to read the GM% from.
    """
    if not _HAS_OPENPYXL:
        return ""
    for pid in phase_ids:
        path = excel_dir / out_quoter_xlsx(pid)
        if not path.exists():
            continue
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                wb = load_workbook(path, data_only=True)
            gm = wb.active["T5"].value
            if isinstance(gm, (int, float)) and 0 < gm < 1:
                return f"{gm:.0%}"
        except Exception:  # noqa: BLE001
            logger.debug("Could not read GM%% from %s", path, exc_info=True)
            continue
    return ""


def compute_phase_revenue(phase: Phase, gm_pct: float = _DEFAULT_GM_PCT) -> float:
    """Compute revenue from domain resource_rows: sum(rate × days) / (1-GM)."""
    month_keys = [f"M{i}" for i in range(1, 7)]
    total_cost = 0.0
    for rr in phase.resource_rows:
        rate = parse_rate_day(
            rr.get(COL_RATE_DAY, ""),
            context=(
                f"phase {phase.phase_id}, stream {rr.get('Stream', '')!r}, "
                f"profile {rr.get('Profile', '')!r}"
            ),
        )
        days = 0.0
        for mk in month_keys:
            with contextlib.suppress(ValueError, TypeError):
                days += float(rr.get(mk, 0) or 0)
        total_cost += rate * days
    if 0 < gm_pct < 1:
        return total_cost / (1 - gm_pct)
    return total_cost
