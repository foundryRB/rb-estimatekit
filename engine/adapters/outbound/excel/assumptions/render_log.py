"""Render Sheet 1: Assumptions Log."""

from __future__ import annotations

from typing import TYPE_CHECKING

from ..branding import RB
from .models import LogData

if TYPE_CHECKING:
    from openpyxl import Workbook


def render_log(wb: Workbook, data: LogData) -> None:
    """Render Sheet 1: Assumptions Log."""
    ws = wb.active
    ws.title = "Assumptions Log"

    headers = ["ID", "Group", "Assumption", "Confidence", "Impact", "Linked Requirements"]
    col_widths = {1: 10, 2: 20, 3: 50, 4: 14, 5: 14, 6: 30}
    RB.set_col_widths(ws, col_widths)

    total_assumptions = sum(len(v) for v in data.grouped_rows.values())
    title = f"{data.project_name} — Assumptions Register"
    subtitle = f"{total_assumptions} assumptions across {len(data.groups)} groups"
    row = RB.apply_title_row(ws, 1, len(headers), title, subtitle)
    row += 1  # spacer

    # ── Confidence legend ───────────────────────────────────
    if data.conf_defs:
        cell = ws.cell(row=row, column=1, value="Confidence Levels")
        cell.font = RB.FONT_BODY_BOLD
        row += 1
        for d in data.conf_defs:
            cell = ws.cell(row=row, column=1, value=d.level)
            cell.font = RB.conf_font(d.level.upper())
            cell.fill = RB.conf_fill(d.level.upper())
            cell.border = RB.BORDER_GRID
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(headers))
            cell = ws.cell(row=row, column=2, value=d.description)
            cell.font = RB.FONT_BODY
            cell.alignment = RB.ALIGN_LEFT
            row += 1
        row += 1

    # ── Impact legend ───────────────────────────────────────
    if data.impact_defs:
        cell = ws.cell(row=row, column=1, value="Impact Levels")
        cell.font = RB.FONT_BODY_BOLD
        row += 1
        for d in data.impact_defs:
            cell = ws.cell(row=row, column=1, value=d.level)
            cell.font = RB.impact_font(d.level.upper())
            cell.fill = RB.impact_fill(d.level.upper())
            cell.border = RB.BORDER_GRID
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(headers))
            cell = ws.cell(row=row, column=2, value=d.description)
            cell.font = RB.FONT_BODY
            cell.alignment = RB.ALIGN_LEFT
            row += 1
        row += 1

    # ── Header ──────────────────────────────────────────────
    row = RB.apply_header_row(ws, row, headers)

    # ── Grouped rows ────────────────────────────────────────
    for cat_idx, (group_name, items) in enumerate(data.grouped_rows.items()):
        # Category separator
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cell = ws.cell(row=row, column=1, value=group_name)
        cell.font = RB.FONT_GROUP
        fill = RB.cat_fill(cat_idx)
        cell.fill = fill
        cell.alignment = RB.ALIGN_LEFT
        for c in range(2, len(headers) + 1):
            ws.cell(row=row, column=c).fill = fill
        ws.row_dimensions[row].height = 22
        row += 1

        for i, a in enumerate(items):
            base_fill = RB.FILL_ALT_ROW if i % 2 == 0 else RB.FILL_WHITE

            vals = [a.assumption_id, a.group, a.assumption, a.confidence, a.impact, a.req_ids_str]
            for c, val in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.border = RB.BORDER_GRID
                cell.alignment = RB.ALIGN_WRAP if c == 3 else RB.ALIGN_LEFT

                if c == 4:  # Confidence
                    cell.font = RB.conf_font(a.confidence)
                    cell.fill = RB.conf_fill(a.confidence)
                elif c == 5:  # Impact
                    cell.font = RB.impact_font(a.impact)
                    cell.fill = RB.impact_fill(a.impact)
                else:
                    cell.font = RB.FONT_BODY
                    cell.fill = base_fill
            row += 1
