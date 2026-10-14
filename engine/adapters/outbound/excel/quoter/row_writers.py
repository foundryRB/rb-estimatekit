"""Row writers for data, separator, and totals rows."""

from __future__ import annotations

from typing import TYPE_CHECKING

from .models import QuoterResourceRow
from .template_styles import apply_styles

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


def write_data_row(
    ws: Worksheet,
    row: int,
    res: QuoterResourceRow,
    styles: dict[int, dict],
    height: float | None,
) -> None:
    """Write a resource data row replicating template row 7."""
    apply_styles(ws, row, styles, height)

    ws.cell(row=row, column=1).value = res.stream                  # A
    # B (Name) left empty
    ws.cell(row=row, column=3).value = res.profile                 # C
    ws.cell(row=row, column=4).value = res.rate                    # D
    ws.cell(row=row, column=5).value = f"=D{row}/8"               # E

    for j, d in enumerate(res.days):                               # F-K
        ws.cell(row=row, column=6 + j).value = d

    ws.cell(row=row, column=12).value = (                          # L
        f"=SUM(Estimate!$F{row}:$K{row})")

    for j in range(len(res.days)):                                 # M-R
        cl = chr(ord("F") + j)
        ws.cell(row=row, column=13 + j).value = f"=$D{row}*{cl}{row}"

    ws.cell(row=row, column=19).value = f"=SUM(M{row}:R{row})"    # S
    ws.cell(row=row, column=20).value = "=$T$5"                    # T
    ws.cell(row=row, column=21).value = 1                          # U
    ws.cell(row=row, column=22).value = (                          # V
        f"=Estimate!$S{row}*Estimate!$U{row}")
    ws.cell(row=row, column=23).value = (                          # W
        f"=Estimate!$S{row}/(1-Estimate!$T{row})")


def write_separator_row(
    ws: Worksheet, row: int, styles: dict[int, dict], height: float | None,
) -> None:
    """Write an empty separator row using template row 8 styling."""
    apply_styles(ws, row, styles, height)


def write_totals_row(
    ws: Worksheet,
    row: int,
    data_start: int,
    styles: dict[int, dict],
    height: float | None,
) -> None:
    """Write the totals row with SUBTOTAL formulas spanning the data range."""
    apply_styles(ws, row, styles, height)

    last = row - 1
    ws.cell(row=row, column=1).value = "Total"                    # A
    ws.cell(row=row, column=4).value = (                          # D
        f"=SUBTOTAL(101,D{data_start}:D{last})")

    for j in range(6):                                            # F-K
        cl = chr(ord("F") + j)
        ws.cell(row=row, column=6 + j).value = (
            f"=SUBTOTAL(109,{cl}{data_start}:{cl}{last})")

    ws.cell(row=row, column=12).value = (                         # L
        f"=SUBTOTAL(109,L{data_start}:L{last})")

    for j in range(6):                                            # M-R
        cl = chr(ord("M") + j)
        ws.cell(row=row, column=13 + j).value = (
            f"=SUBTOTAL(109,{cl}{data_start}:{cl}{last})")

    ws.cell(row=row, column=19).value = (                         # S
        f"=SUBTOTAL(109,S{data_start}:S{last})")
    ws.cell(row=row, column=22).value = (                         # V
        f"=SUBTOTAL(109,V{data_start}:V{last})")
    ws.cell(row=row, column=23).value = (                         # W
        f"=SUBTOTAL(109,W{data_start}:W{last})")
