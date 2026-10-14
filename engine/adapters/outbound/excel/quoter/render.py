"""Render the quoter workbook from template and prepared data."""

from __future__ import annotations

import warnings
from pathlib import Path

from openpyxl import load_workbook

from .models import QuoterPhaseData
from .row_writers import write_data_row, write_separator_row, write_totals_row
from .template_styles import MAX_COL, snapshot_row


def render_quoter(
    template_path: Path,
    data: QuoterPhaseData,
    output_dir: Path,
) -> str:
    """Load the template, render data, and save. Returns filename."""
    from engine.domain.constants.files.outputs import out_quoter_xlsx

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(template_path)
    ws = wb.active

    # Remove Excel table definitions that break when row count changes
    for name in list(getattr(ws, "tables", {})):
        del ws.tables[name]

    # Update title
    ws["A1"].value = f"{data.project_name} — {data.phase_name}"

    # Snapshot template rows before overwriting
    data_styles, data_height = snapshot_row(ws, 7)
    sep_styles, sep_height = snapshot_row(ws, 8)
    total_styles, total_height = snapshot_row(ws, 9)

    n_output = len(data.plan)

    # Clear cells from row 7 down
    clear_end = max(ws.max_row, 7 + n_output + 10)
    for r in range(7, clear_end + 1):
        for c in range(1, MAX_COL + 1):
            ws.cell(row=r, column=c).value = None

    # Write data / separator rows starting at row 7
    data_start = 7
    for i, (rtype, res) in enumerate(data.plan):
        row = data_start + i
        if rtype == "sep":
            write_separator_row(ws, row, sep_styles, sep_height)
        else:
            assert res is not None
            write_data_row(ws, row, res, data_styles, data_height)

    # Totals row immediately after the last data/separator row
    totals_row = data_start + n_output
    write_totals_row(ws, totals_row, data_start, total_styles, total_height)

    # Rewrite header formulas that referenced the deleted Excel Table
    ws["W4"] = f"=(W{totals_row}-V{totals_row})/W{totals_row}"
    ws["B2"] = f"=SUBTOTAL(109,Estimate!$W$7:$W${totals_row - 1})"

    fname = out_quoter_xlsx(data.phase_id)
    wb.save(output_dir / fname)
    return fname
