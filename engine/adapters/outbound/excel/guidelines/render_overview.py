"""Render Sheet 1: Guidelines Overview."""

from __future__ import annotations

from typing import TYPE_CHECKING

from engine.adapters.outbound.excel.branding import CF

if TYPE_CHECKING:
    from openpyxl import Workbook

from .models import GuidelinesData

__all__ = ["render_overview"]


def render_overview(wb: Workbook, data: GuidelinesData) -> None:
    """Render Sheet 1: Guidelines Overview."""
    ws = wb.active
    ws.title = "Guidelines Overview"

    headers = ["Group", "ID", "Guideline", "Description", "Phase IDs"]
    col_widths = {1: 22, 2: 12, 3: 35, 4: 55, 5: 14}
    CF.set_col_widths(ws, col_widths)

    title = f"{data.project_name} — Architecture Guidelines"
    subtitle = f"{data.total_count} guidelines"
    row = CF.apply_title_row(ws, 1, len(headers), title, subtitle)
    row += 1

    row = CF.apply_header_row(ws, row, headers)

    total_count = 0
    for cat_idx, (group_name, items) in enumerate(data.groups):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cell = ws.cell(row=row, column=1, value=group_name)
        cell.font = CF.FONT_GROUP
        fill = CF.cat_fill(cat_idx)
        cell.fill = fill
        cell.alignment = CF.ALIGN_LEFT
        for c in range(2, len(headers) + 1):
            ws.cell(row=row, column=c).fill = fill
        ws.row_dimensions[row].height = 22
        row += 1

        for i, g in enumerate(items):
            base_fill = CF.FILL_ALT_ROW if i % 2 == 0 else CF.FILL_WHITE
            vals = [g.group, g.guideline_id, g.name, g.description, g.phase_ids]
            for c, val in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.font = CF.FONT_BODY
                cell.fill = base_fill
                cell.border = CF.BORDER_GRID
                cell.alignment = CF.ALIGN_WRAP if c == 4 else CF.ALIGN_LEFT
            total_count += 1
            row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    cell = ws.cell(row=row, column=1, value="TOTAL GUIDELINES")
    cell.font = CF.FONT_TOTAL
    cell.fill = CF.FILL_TOTAL
    cell.alignment = CF.ALIGN_RIGHT
    for c in range(2, 4):
        ws.cell(row=row, column=c).fill = CF.FILL_TOTAL

    cell = ws.cell(row=row, column=4, value=total_count)
    cell.font = CF.FONT_TOTAL
    cell.fill = CF.FILL_TOTAL
    cell.alignment = CF.ALIGN_CENTER
    cell.border = CF.BORDER_GRID
    ws.row_dimensions[row].height = 26
