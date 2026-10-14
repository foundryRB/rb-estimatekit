"""Render Sheet 2: Summary by Group."""

from __future__ import annotations

from typing import TYPE_CHECKING

from engine.adapters.outbound.excel.branding import CF

if TYPE_CHECKING:
    from openpyxl import Workbook

from .models import GuidelinesData

__all__ = ["render_summary"]


def render_summary(wb: Workbook, data: GuidelinesData) -> None:
    """Render Sheet 2: Summary by Group."""
    ws = wb.create_sheet("Summary by Group")

    headers = ["Group", "Count", "Guidelines"]
    col_widths = {1: 25, 2: 10, 3: 60}
    CF.set_col_widths(ws, col_widths)

    row = CF.apply_title_row(
        ws, 1, len(headers),
        f"{data.project_name} — Guidelines Summary",
    )
    row += 1
    row = CF.apply_header_row(ws, row, headers)

    for i, s in enumerate(data.summaries):
        fill = CF.FILL_ALT_ROW if i % 2 == 0 else CF.FILL_WHITE

        cell = ws.cell(row=row, column=1, value=s.group)
        cell.font = CF.FONT_BODY_BOLD
        cell.fill = fill
        cell.border = CF.BORDER_GRID
        cell.alignment = CF.ALIGN_LEFT

        cell = ws.cell(row=row, column=2, value=s.count)
        cell.font = CF.FONT_BODY
        cell.fill = fill
        cell.border = CF.BORDER_GRID
        cell.alignment = CF.ALIGN_CENTER

        cell = ws.cell(row=row, column=3, value=s.names)
        cell.font = CF.FONT_BODY
        cell.fill = fill
        cell.border = CF.BORDER_GRID
        cell.alignment = CF.ALIGN_WRAP

        row += 1
