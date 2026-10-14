"""Workbook styles, shared by the Excel generators.

CF owns the installed application's spreadsheet palette and typography.
Semantic risk and category colours retain their data meaning.
"""

from __future__ import annotations

from typing import ClassVar

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

from engine.domain.constants.colors import PHASE_COLORS_HEX

__all__ = ["CF"]


class CF:
    """Brand constants for Excel formatting."""

    BRAND = "00A86B"
    BLACK = "000000"
    DARK = "1A1A1A"
    DARK_GREY = "333333"
    MID_GREY = "666666"
    LIGHT_GREY = "E8E8E8"
    BG_CARD = "F2F2F2"
    BG_SUBTLE = "F9F9F9"
    WHITE = "FFFFFF"

    PHASE_COLORS: ClassVar[dict[str, str]] = PHASE_COLORS_HEX

    CAT_COLORS: ClassVar[list[str]] = [
        "F9E4E4", "FFF8E1", "E8F5E9", "E3F2FD", "F3E5F5",
        "FFF3E0", "E0F7FA", "FCE4EC", "F1F8E9", "EFEBE9",
    ]

    CONF_COLORS: ClassVar[dict[str, tuple[str, str]]] = {
        "HIGH": ("2E7D32", "E8F5E9"),
        "MEDIUM": ("E65100", "FFF8E1"),
        "LOW": ("C62828", "FCE4EC"),
    }

    IMPACT_COLORS: ClassVar[dict[str, tuple[str, str]]] = {
        "CRITICAL": ("FFFFFF", "C62828"),
        "HIGH": ("C62828", "FCE4EC"),
        "MEDIUM": ("E65100", "FFF8E1"),
        "LOW": ("2E7D32", "E8F5E9"),
    }

    HEAT_COLORS: ClassVar[dict[tuple[str, str], str]] = {
        ("HIGH", "CRITICAL"): "C62828",
        ("HIGH", "HIGH"): "E57373",
        ("HIGH", "MEDIUM"): "FFF8E1",
        ("HIGH", "LOW"): "E8F5E9",
        ("MEDIUM", "CRITICAL"): "C62828",
        ("MEDIUM", "HIGH"): "E57373",
        ("MEDIUM", "MEDIUM"): "FFF8E1",
        ("MEDIUM", "LOW"): "E8F5E9",
        ("LOW", "CRITICAL"): "C62828",
        ("LOW", "HIGH"): "C62828",
        ("LOW", "MEDIUM"): "E57373",
        ("LOW", "LOW"): "FFF8E1",
    }

    FONT_FAMILY = "Segoe UI"

    FONT_TITLE = Font(name=FONT_FAMILY, size=16, bold=True, color=WHITE)
    FONT_SUBTITLE = Font(name=FONT_FAMILY, size=11, italic=True, color=LIGHT_GREY)
    FONT_HEADER = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
    FONT_BODY = Font(name=FONT_FAMILY, size=10, color=DARK)
    FONT_BODY_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color=DARK)
    FONT_SMALL = Font(name=FONT_FAMILY, size=9, color=MID_GREY)
    FONT_CHECK = Font(name=FONT_FAMILY, size=12, color=DARK)
    FONT_GROUP = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
    FONT_TOTAL = Font(name=FONT_FAMILY, size=11, bold=True, color=WHITE)
    FONT_NUMBER = Font(name=FONT_FAMILY, size=10, color=DARK)
    FONT_PARTIAL = Font(name=FONT_FAMILY, size=9, italic=True, color=MID_GREY)

    FILL_TITLE = PatternFill("solid", fgColor=DARK)
    FILL_HEADER = PatternFill("solid", fgColor=DARK)
    FILL_BRAND = PatternFill("solid", fgColor=BRAND)
    FILL_ALT_ROW = PatternFill("solid", fgColor=BG_SUBTLE)
    FILL_CARD = PatternFill("solid", fgColor=BG_CARD)
    FILL_LIGHT_GREY = PatternFill("solid", fgColor=LIGHT_GREY)
    FILL_GROUP = PatternFill("solid", fgColor=MID_GREY)
    FILL_TOTAL = PatternFill("solid", fgColor=BRAND)
    FILL_WHITE = PatternFill("solid", fgColor=WHITE)
    FILL_NONE = PatternFill(fill_type=None)

    ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

    THIN = Side(style="thin", color=LIGHT_GREY)
    HAIR = Side(style="hair", color=LIGHT_GREY)
    BORDER_GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    BORDER_BOTTOM = Border(bottom=Side(style="thin", color=MID_GREY))
    BORDER_NONE = Border()

    @staticmethod
    def conf_font(level: str) -> Font:
        """Font for confidence-level cells."""
        fg, _ = CF.CONF_COLORS.get(level, (CF.DARK, CF.WHITE))
        return Font(name=CF.FONT_FAMILY, size=10, bold=True, color=fg)

    @staticmethod
    def conf_fill(level: str) -> PatternFill:
        """Fill for confidence-level cells."""
        _, bg = CF.CONF_COLORS.get(level, (CF.DARK, CF.WHITE))
        return PatternFill("solid", fgColor=bg)

    @staticmethod
    def impact_font(level: str) -> Font:
        """Font for impact-level cells."""
        fg, _ = CF.IMPACT_COLORS.get(level, (CF.DARK, CF.WHITE))
        return Font(name=CF.FONT_FAMILY, size=10, bold=True, color=fg)

    @staticmethod
    def impact_fill(level: str) -> PatternFill:
        """Fill for impact-level cells."""
        _, bg = CF.IMPACT_COLORS.get(level, (CF.DARK, CF.WHITE))
        return PatternFill("solid", fgColor=bg)

    @staticmethod
    def heat_fill(conf: str, impact: str) -> PatternFill:
        """Fill for a risk-matrix heat-map cell."""
        color = CF.HEAT_COLORS.get((conf, impact), CF.WHITE)
        return PatternFill("solid", fgColor=color)

    @staticmethod
    def heat_font(conf: str, impact: str) -> Font:
        """Font for a risk-matrix heat-map cell."""
        heat = CF.HEAT_COLORS.get((conf, impact), CF.WHITE)
        fg = "FFFFFF" if heat in ("C62828",) else CF.DARK
        return Font(name=CF.FONT_FAMILY, size=10, bold=True, color=fg)

    @staticmethod
    def phase_fill(phase_id: str) -> PatternFill:
        color = CF.PHASE_COLORS.get(phase_id, CF.DARK_GREY)
        return PatternFill("solid", fgColor=color)

    @staticmethod
    def cat_fill(idx: int) -> PatternFill:
        color = CF.CAT_COLORS[idx % len(CF.CAT_COLORS)]
        return PatternFill("solid", fgColor=color)

    @staticmethod
    def apply_title_row(ws, row: int, max_col: int, title: str, subtitle: str | None = None) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = CF.FONT_TITLE
        cell.fill = CF.FILL_TITLE
        cell.alignment = CF.ALIGN_LEFT
        ws.row_dimensions[row].height = 32
        if subtitle:
            ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=max_col)
            cell = ws.cell(row=row + 1, column=1, value=subtitle)
            cell.font = CF.FONT_SUBTITLE
            cell.fill = CF.FILL_TITLE
            cell.alignment = CF.ALIGN_LEFT
            ws.row_dimensions[row + 1].height = 20
            return row + 2
        return row + 1

    @staticmethod
    def apply_header_row(ws, row: int, headers: list[str], col_start: int = 1) -> int:
        for i, h in enumerate(headers):
            cell = ws.cell(row=row, column=col_start + i, value=h)
            cell.font = CF.FONT_HEADER
            cell.fill = CF.FILL_HEADER
            cell.alignment = CF.ALIGN_CENTER
            cell.border = CF.BORDER_GRID
        ws.row_dimensions[row].height = 24
        return row + 1

    @staticmethod
    def set_col_widths(ws, widths: dict) -> None:
        from openpyxl.utils import get_column_letter
        for key, w in widths.items():
            if isinstance(key, int):
                key = get_column_letter(key)
            ws.column_dimensions[key].width = w
