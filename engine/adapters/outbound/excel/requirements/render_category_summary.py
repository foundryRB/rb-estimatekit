"""Render Sheet 3: Category Summary."""

from __future__ import annotations

from typing import TYPE_CHECKING

from ..branding import CF
from .models import RequirementsData

if TYPE_CHECKING:
    from openpyxl import Workbook


def render_category_summary(wb: Workbook, data: RequirementsData) -> None:
    """Render Sheet 3: Category Summary."""
    ws = wb.create_sheet("Category Summary")
    ws.sheet_properties.tabColor = "999999"

    num_phases = len(data.phase_ids)
    total_cols = 2 + num_phases

    widths: dict[int, int] = {1: 40, 2: 14}
    for i in range(num_phases):
        widths[3 + i] = 14
    CF.set_col_widths(ws, widths)

    # Title
    row = 1
    row = CF.apply_title_row(
        ws,
        row,
        total_cols,
        "Requirements by Category — Phased View",
        "Summarized view showing how each category is fulfilled across phases",
    )
    row += 1

    # Header
    headers = ["Category", "Total Reqs"] + data.phase_names
    row = CF.apply_header_row(ws, row, headers)

    # Data rows
    for cat_idx, cs in enumerate(data.category_summaries):
        fill = CF.cat_fill(cat_idx)

        ws.cell(row=row, column=1, value=cs.display_name).font = CF.FONT_BODY_BOLD
        ws.cell(row=row, column=2, value=cs.total).font = CF.FONT_BODY_BOLD

        for p_idx, pid in enumerate(data.phase_ids):
            col = 3 + p_idx
            val = cs.phase_counts.get(pid, 0)
            cell = ws.cell(row=row, column=col)
            if val == 0:
                cell.value = "—"
                cell.font = CF.FONT_PARTIAL
            else:
                cell.value = val
                cell.font = CF.FONT_BODY
            cell.alignment = CF.ALIGN_CENTER

        for c in range(1, total_cols + 1):
            ws.cell(row=row, column=c).fill = fill
            ws.cell(row=row, column=c).border = CF.BORDER_GRID
        ws.row_dimensions[row].height = 24
        row += 1

    # Total row
    ws.cell(row=row, column=1, value="TOTAL").font = CF.FONT_TOTAL
    ws.cell(row=row, column=2, value=data.total_reqs).font = CF.FONT_TOTAL

    for p_idx, pid in enumerate(data.phase_ids):
        col = 3 + p_idx
        phase_total = sum(cs.phase_counts.get(pid, 0) for cs in data.category_summaries)
        ws.cell(row=row, column=col, value=phase_total).font = CF.FONT_TOTAL
        ws.cell(row=row, column=col).alignment = CF.ALIGN_CENTER

    for c in range(1, total_cols + 1):
        ws.cell(row=row, column=c).fill = CF.FILL_TOTAL
        ws.cell(row=row, column=c).alignment = CF.ALIGN_CENTER
        ws.cell(row=row, column=c).border = CF.BORDER_GRID
    ws.row_dimensions[row].height = 26
