"""Render Sheet 2: Requirements by Phase."""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.styles import Font

from ..branding import RB
from .models import RequirementsData

if TYPE_CHECKING:
    from openpyxl import Workbook


def render_requirements(wb: Workbook, data: RequirementsData) -> None:
    """Render Sheet 2: Requirements by Phase."""
    ws = wb.create_sheet("Requirements by Phase")
    ws.sheet_properties.tabColor = RB.DARK_GREY

    num_phases = len(data.phase_ids)
    total_cols = 2 + num_phases

    # Column widths
    widths: dict[int, int] = {1: 14, 2: 55}
    for i in range(num_phases):
        widths[3 + i] = 14
    RB.set_col_widths(ws, widths)

    # Title
    row = 1
    row = RB.apply_title_row(
        ws,
        row,
        total_cols,
        "Complete Requirements Matrix",
        f"{data.total_reqs} requirements mapped to delivery phases and categories",
    )
    row += 1  # blank

    # Header
    headers = ["Req ID", "Requirement Description"] + data.phase_names
    row = RB.apply_header_row(ws, row, headers)

    # Group by category
    for cat_idx, (cat_display, cat_rows) in enumerate(data.req_groups):
        # Category separator row
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=total_cols
        )
        cell = ws.cell(row=row, column=1, value=cat_display)
        cell.font = RB.FONT_GROUP
        cell.fill = RB.FILL_GROUP
        cell.alignment = RB.ALIGN_LEFT
        ws.row_dimensions[row].height = 22
        row += 1

        fill = RB.cat_fill(cat_idx)

        for req in cat_rows:
            ws.cell(row=row, column=1, value=req.req_id).font = RB.FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=req.description).font = RB.FONT_BODY

            for p_idx, pid in enumerate(data.phase_ids):
                col = 3 + p_idx
                alloc = req.phase_allocations.get(pid, "")
                cell = ws.cell(row=row, column=col)
                if alloc == "✓":
                    cell.value = "✓"
                    cell.font = Font(
                        name=RB.FONT_FAMILY, size=12, bold=True, color="2E7D32"
                    )
                elif alloc and alloc not in ("—", "-"):
                    cell.value = alloc
                    cell.font = RB.FONT_PARTIAL
                cell.alignment = RB.ALIGN_CENTER

            for c in range(1, total_cols + 1):
                ws.cell(row=row, column=c).fill = fill
                ws.cell(row=row, column=c).border = RB.BORDER_GRID
                if c <= 2:
                    ws.cell(row=row, column=c).alignment = RB.ALIGN_LEFT
            ws.row_dimensions[row].height = 22
            row += 1
