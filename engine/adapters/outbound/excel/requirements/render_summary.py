"""Render Sheet 1: Summary."""

from __future__ import annotations

from typing import TYPE_CHECKING

from ..branding import RB
from .models import RequirementsData

if TYPE_CHECKING:
    from openpyxl import Workbook


def render_summary(wb: Workbook, data: RequirementsData) -> None:
    """Render Sheet 1: Summary."""
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = RB.BRAND

    # Column widths
    RB.set_col_widths(ws, {1: 28, 2: 18, 3: 55})

    # Title
    row = 1
    row = RB.apply_title_row(
        ws,
        row,
        3,
        f"{data.project_name} — Customer Requirements",
        data.project_description,
    )
    row += 1  # blank

    # Phase overview header
    headers = ["Phase", "Requirements", "Key Capabilities"]
    row = RB.apply_header_row(ws, row, headers)

    # Phase rows
    for i, po in enumerate(data.phase_overview):
        fill = RB.cat_fill(i)

        ws.cell(row=row, column=1, value=po.name).font = RB.FONT_BODY_BOLD
        ws.cell(row=row, column=2, value=po.req_label)
        ws.cell(row=row, column=3, value=po.description)

        for c in range(1, 4):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill
            if not cell.font.bold:
                cell.font = RB.FONT_BODY
            cell.alignment = RB.ALIGN_LEFT
            cell.border = RB.BORDER_GRID
        ws.row_dimensions[row].height = 28
        row += 1

    row += 1  # blank

    # Investment Summary
    max_col = 3
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    cell = ws.cell(row=row, column=1, value="INVESTMENT SUMMARY")
    cell.font = RB.FONT_HEADER
    cell.fill = RB.FILL_BRAND
    cell.alignment = RB.ALIGN_LEFT
    ws.row_dimensions[row].height = 24
    row += 1

    # Effort sub-header
    effort_headers = ["Phase", "Build (MD)", "Overhead (MD)", "Total (MD)"]
    RB.set_col_widths(ws, {1: 28, 2: 18, 3: 18, 4: 18})
    max_col = len(effort_headers)
    row = RB.apply_header_row(ws, row, effort_headers)

    for er in data.effort_rows:
        ws.cell(row=row, column=1, value=er.name).font = RB.FONT_BODY
        ws.cell(row=row, column=2, value=er.build).font = RB.FONT_NUMBER
        ws.cell(row=row, column=3, value=er.overhead).font = RB.FONT_NUMBER
        ws.cell(row=row, column=4, value=er.total).font = RB.FONT_NUMBER

        for c in range(1, max_col + 1):
            ws.cell(row=row, column=c).border = RB.BORDER_GRID
            ws.cell(row=row, column=c).alignment = RB.ALIGN_CENTER
        ws.cell(row=row, column=1).alignment = RB.ALIGN_LEFT
        row += 1

    # Total row
    ws.cell(row=row, column=1, value="TOTAL").font = RB.FONT_TOTAL
    ws.cell(row=row, column=2, value=data.grand_build).font = RB.FONT_TOTAL
    ws.cell(row=row, column=3, value=data.grand_overhead).font = RB.FONT_TOTAL
    ws.cell(row=row, column=4, value=data.grand_total).font = RB.FONT_TOTAL

    for c in range(1, max_col + 1):
        ws.cell(row=row, column=c).fill = RB.FILL_TOTAL
        ws.cell(row=row, column=c).alignment = RB.ALIGN_CENTER
        ws.cell(row=row, column=c).border = RB.BORDER_GRID
    ws.cell(row=row, column=1).alignment = RB.ALIGN_LEFT
    ws.row_dimensions[row].height = 26
