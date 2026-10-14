"""Render the Financial Model workbook from prepared data."""

from __future__ import annotations

import warnings
from pathlib import Path

from openpyxl import load_workbook

from engine.domain.constants.files.outputs import out_financial_xlsx

from .models import FinancialData

__all__ = ["render_financial"]


def render_financial(
    template_path: Path,
    data: FinancialData,
    output_dir: Path,
) -> str:
    """Load the template, apply placeholders, and save. Returns filename."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(template_path)
    ws = wb.active

    ws["A1"].value = f"Financial Model: {data.phase_name}"

    for row in ws.iter_rows():
        for cell in row:
            if (
                cell.value
                and isinstance(cell.value, str)
                and "{{totalStaffDays}}" in cell.value
            ):
                cell.value = data.total_staff_days

    fname = out_financial_xlsx(data.phase_id)
    wb.save(output_dir / fname)
    return fname
