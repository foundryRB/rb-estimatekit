"""Render Sheet 1: Build Activities."""

from __future__ import annotations

from typing import TYPE_CHECKING

from ..branding import CF
from .models import EstimationData

if TYPE_CHECKING:
    from openpyxl import Workbook


def render_build(wb: Workbook, data: EstimationData) -> None:
    """Render Sheet 1: Build Activities."""
    ws = wb.active
    ws.title = "Build Activities"

    headers = ["Activity", "Linked Requirements", "Effort (MD)", "Description"]
    col_widths = {1: 40, 2: 30, 3: 14, 4: 40}
    CF.set_col_widths(ws, col_widths)

    row = CF.apply_title_row(ws, 1, len(headers), f"{data.phase_name} — Build Activities")
    row += 1  # spacer

    row = CF.apply_header_row(ws, row, headers)

    phase_fill = CF.phase_fill(data.phase_id)
    for group_idx, group in enumerate(data.build_groups, start=1):
        # Group separator row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cell = ws.cell(row=row, column=1, value=f"GROUP {group_idx}: {group.group_name}")
        cell.font = CF.FONT_GROUP
        cell.fill = CF.FILL_GROUP
        cell.alignment = CF.ALIGN_LEFT
        for c in range(2, len(headers) + 1):
            ws.cell(row=row, column=c).fill = CF.FILL_GROUP
        ws.row_dimensions[row].height = 22
        row += 1

        for i, act in enumerate(group.activities):
            fill = CF.FILL_ALT_ROW if i % 2 == 0 else CF.FILL_WHITE
            vals = [act.name, act.req_ids_str, act.effort_md, ""]
            for c, val in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.font = CF.FONT_BODY
                cell.fill = fill
                cell.border = CF.BORDER_GRID
                cell.alignment = CF.ALIGN_RIGHT if c == 3 else CF.ALIGN_LEFT
            row += 1

        # Group subtotal
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row=row, column=1, value=f"{group.group_name} Subtotal")
        cell.font = CF.FONT_BODY_BOLD
        cell.fill = CF.FILL_CARD
        cell.alignment = CF.ALIGN_RIGHT
        ws.cell(row=row, column=2).fill = CF.FILL_CARD

        cell = ws.cell(row=row, column=3, value=group.subtotal)
        cell.font = CF.FONT_BODY_BOLD
        cell.fill = CF.FILL_CARD
        cell.alignment = CF.ALIGN_RIGHT
        cell.border = CF.BORDER_GRID

        ws.cell(row=row, column=4, value="").fill = CF.FILL_CARD
        row += 1

    # BUILD TOTAL
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=1, value="BUILD TOTAL")
    cell.font = CF.FONT_TOTAL
    cell.fill = phase_fill
    cell.alignment = CF.ALIGN_RIGHT
    ws.cell(row=row, column=2).fill = phase_fill

    cell = ws.cell(row=row, column=3, value=data.build_total)
    cell.font = CF.FONT_TOTAL
    cell.fill = phase_fill
    cell.alignment = CF.ALIGN_RIGHT
    cell.border = CF.BORDER_GRID

    ws.cell(row=row, column=4, value="").fill = phase_fill
    ws.row_dimensions[row].height = 26
